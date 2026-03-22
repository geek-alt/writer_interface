import asyncio
import re
from pathlib import Path

from docx import Document
from ebooklib import epub
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from backend.core.state_manager import StateManager


class NovelExporter:
    def __init__(self, state_manager: StateManager) -> None:
        self.state = state_manager

    async def export_docx(self, project_id: str, chapters: list[dict]) -> Path:
        return await asyncio.to_thread(self._export_docx_sync, project_id, chapters)

    async def export_epub(self, project_id: str, chapters: list[dict]) -> Path:
        return await asyncio.to_thread(self._export_epub_sync, project_id, chapters)

    async def export_pdf(self, project_id: str, chapters: list[dict]) -> Path:
        return await asyncio.to_thread(self._export_pdf_sync, project_id, chapters)

    def _exports_dir(self, project_id: str) -> Path:
        export_dir = Path(self.state.PROJECTS_DIR) / project_id / "exports"
        export_dir.mkdir(parents=True, exist_ok=True)
        return export_dir

    async def _project_title(self, project_id: str) -> str:
        config = await self.state.load_config(project_id)
        return config.get("project_name") or f"NovelForge_{project_id}"

    def _safe_stem(self, text: str) -> str:
        cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", text).strip("_")
        return cleaned or "novelforge_export"

    def _chapter_title(self, chapter: dict) -> str:
        number = chapter.get("number", "?")
        title = chapter.get("title", f"Chapter {number}")
        return f"Chapter {number}: {title}"

    def _export_docx_sync(self, project_id: str, chapters: list[dict]) -> Path:
        title = chapters[0].get("project_name", "NovelForge Export") if chapters else "NovelForge Export"
        doc = Document()
        doc.add_heading(title, level=0)

        for chapter in chapters:
            doc.add_page_break()
            doc.add_heading(self._chapter_title(chapter), level=1)
            for para in str(chapter.get("content", "")).split("\n\n"):
                p = para.strip()
                if p:
                    doc.add_paragraph(p)

        file_path = self._exports_dir(project_id) / f"{self._safe_stem(title)}.docx"
        doc.save(file_path)
        return file_path

    def _export_epub_sync(self, project_id: str, chapters: list[dict]) -> Path:
        book = epub.EpubBook()
        title = chapters[0].get("project_name", "NovelForge Export") if chapters else "NovelForge Export"
        identifier = self._safe_stem(f"{project_id}_{title}")

        book.set_identifier(identifier)
        book.set_title(title)
        book.set_language("en")
        book.add_author("NovelForge")

        toc_items = []
        spine_items = ["nav"]

        for chapter in chapters:
            ch_number = chapter.get("number", 0)
            ch_title = self._chapter_title(chapter)
            content_html = "".join(
                f"<p>{self._escape_html(p.strip())}</p>"
                for p in str(chapter.get("content", "")).split("\n\n")
                if p.strip()
            )

            epub_chapter = epub.EpubHtml(
                title=ch_title,
                file_name=f"chapter_{int(ch_number):03d}.xhtml" if isinstance(ch_number, int) else f"chapter_{ch_number}.xhtml",
                lang="en",
            )
            epub_chapter.content = f"<h1>{self._escape_html(ch_title)}</h1>{content_html}"
            book.add_item(epub_chapter)
            toc_items.append(epub_chapter)
            spine_items.append(epub_chapter)

        book.toc = tuple(toc_items)
        book.spine = spine_items
        book.add_item(epub.EpubNcx())
        book.add_item(epub.EpubNav())

        file_path = self._exports_dir(project_id) / f"{self._safe_stem(title)}.epub"
        epub.write_epub(str(file_path), book, {})
        return file_path

    def _export_pdf_sync(self, project_id: str, chapters: list[dict]) -> Path:
        title = chapters[0].get("project_name", "NovelForge Export") if chapters else "NovelForge Export"
        file_path = self._exports_dir(project_id) / f"{self._safe_stem(title)}.pdf"

        styles = getSampleStyleSheet()
        heading_style = styles["Heading1"]
        body_style = styles["BodyText"]

        story = [Paragraph(title, styles["Title"]), Spacer(1, 18)]

        for chapter in chapters:
            story.append(Paragraph(self._chapter_title(chapter), heading_style))
            story.append(Spacer(1, 10))
            for para in str(chapter.get("content", "")).split("\n\n"):
                p = para.strip()
                if p:
                    story.append(Paragraph(self._escape_html(p), body_style))
                    story.append(Spacer(1, 8))
            story.append(Spacer(1, 18))

        doc = SimpleDocTemplate(str(file_path), pagesize=LETTER)
        doc.build(story)
        return file_path

    def _escape_html(self, text: str) -> str:
        return (
            text.replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )
